"""재고 대량 가져오기 (M050).

한 엔드포인트에 단계를 넘긴다
--------------------------
`/analyze` `/dry-run` `/apply` 를 따로 두지 않고 `mode` 를 받는다. 세 단계가
**같은 파일에 대해** 도는 것이 중요하고, 경로를 나누면 사용자가 분석은 A 파일로
하고 적용은 B 파일로 하는 일이 생긴다.

기본값은 `DRY_RUN` 이다. `APPLY` 는 명시적으로 보내야 한다 — 실수로 1,000행이
반영되는 경로를 만들지 않는다.

파일을 저장하지 않는다
-------------------
업로드된 내용은 메모리에서 읽고 버린다. 해시만 남는다. 재고 파일에는 상품
목록이 통째로 들어 있어 보관할 이유가 없고, 보관하면 지워야 할 것이 하나 더
생긴다.
"""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import date
from typing import Any

from fastapi import APIRouter, Depends, File, Form, Query, UploadFile
from sqlalchemy.orm import Session

from app.core.exceptions import ValidationError
from app.core.rbac import require_roles
from app.core.response import envelope
from app.deps import CurrentUser, get_db
from app.services import inventory_import as imp

router = APIRouter(prefix="/inventory/import", tags=["inventory"])

#: 재고를 바꾸는 일이라 조정과 같은 권한이다.
_manage = require_roles("ADMIN", "OPERATOR")
_view = require_roles("ADMIN", "OPERATOR", "SOURCING", "ACCOUNTING")

#: 업로드 상한. 1,000행짜리 재고 파일은 수백 KB 다. 10MB 면 충분하고,
#: 이보다 크면 잘못된 파일일 가능성이 높다.
_MAX_BYTES = 10 * 1024 * 1024


def _rows_from_upload(name: str, raw: bytes) -> list[tuple[Any, ...]]:
    """CSV 와 XLSX 를 같은 모양(행들의 튜플)으로 만든다.

    CSV 를 함께 받는 이유: 엑셀이 없는 환경에서도 만들 수 있고, 문제가 생겼을
    때 사용자가 내용을 직접 눈으로 확인할 수 있다.
    """
    lower = name.lower()
    if lower.endswith(".csv"):
        # BOM 이 붙은 CSV 가 흔하다(엑셀에서 저장하면 붙는다). 안 벗기면
        # 첫 열 이름이 안 읽힌다.
        text_data = raw.decode("utf-8-sig", errors="replace")
        return [tuple(r) for r in csv.reader(io.StringIO(text_data))]

    if lower.endswith((".xlsx", ".xlsm")):
        import openpyxl

        wb = openpyxl.load_workbook(
            io.BytesIO(raw), data_only=True, read_only=True)
        try:
            sheet = wb[wb.sheetnames[0]]
            return list(sheet.iter_rows(values_only=True))
        finally:
            wb.close()

    raise ValidationError(
        "CSV 또는 XLSX 파일만 받습니다. 확장자를 확인하세요.")


@router.post("")
async def run_import(
    file: UploadFile = File(...),
    # 기본이 DRY_RUN 이다. 실수로 반영되지 않게 한다.
    mode: str = Form("DRY_RUN"),
    as_of_date: date | None = Form(None),
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(_manage),
) -> dict:
    """재고 파일 분석·검토·반영.

    ```
    ANALYZE   파일만 읽는다 (DB 를 보지 않는다)
    DRY_RUN   무엇이 바뀔지 계산한다 (쓰지 않는다)
    APPLY     원장에 반영한다
    ```
    """
    if mode not in imp.MODES:
        raise ValidationError(f"알 수 없는 단계입니다: {mode}")

    raw = await file.read()
    if not raw:
        raise ValidationError("파일이 비어 있습니다.")
    if len(raw) > _MAX_BYTES:
        raise ValidationError("파일이 너무 큽니다 (최대 10MB).")

    # 해시는 **내용**에서 만든다. 파일명은 바뀌고 엑셀은 열기만 해도
    # 수정 시각이 바뀐다.
    sha256 = hashlib.sha256(raw).hexdigest()
    rows = _rows_from_upload(file.filename or "upload", raw)

    plan = imp.plan(
        db, org_id=user.org_id, rows=rows, mode=mode, as_of_date=as_of_date,
        source_sha256=sha256)

    if mode != "APPLY":
        return envelope({**plan.as_dict(), "source_name": file.filename})

    return envelope(imp.apply_plan(
        db, org_id=user.org_id, actor_user_id=user.user_id, plan_result=plan,
        source_name=file.filename or "upload", source_sha256=sha256))


@router.get("/runs")
def list_runs(
    limit: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(_view),
) -> dict:
    """지난 가져오기 이력. 무엇을 언제 넣었는지 되짚는 화면이다."""
    return envelope(imp.list_runs(db, org_id=user.org_id, limit=limit))


@router.get("/template")
def template(user: CurrentUser = Depends(_view)) -> dict:
    """어떤 열이 필요한지. 화면이 이걸로 안내 문구와 예시를 만든다.

    파일을 내려주지 않고 **형식만** 알려준다 — 사용자는 대개 자기 파일을
    이미 갖고 있고, 우리 양식으로 옮겨 적게 하면 오히려 오류가 는다.
    """
    _ = user
    return envelope({
        "identity_columns": list(imp.IDENTITY_COLUMNS),
        "columns": [
            {"field": "sku", "label": "SKU", "required": False,
             "note": "상품 코드. 셋 중 하나는 있어야 합니다.",
             "aliases": ["SKU", "상품코드", "품번"]},
            {"field": "business_key", "label": "상품번호", "required": False,
             "note": "PROD-000000 형식", "aliases": ["상품번호", "관리번호"]},
            {"field": "product_id", "label": "상품 ID", "required": False,
             "note": "UUID", "aliases": ["product_id"]},
            {"field": "variant_key", "label": "옵션코드", "required": False,
             "note": "옵션이 있는 상품은 필수입니다.",
             "aliases": ["옵션코드", "옵션"]},
            {"field": "quantity", "label": "수량", "required": True,
             "note": "정수. 0 도 유효한 값입니다.",
             "aliases": ["수량", "재고", "실사수량"]},
            {"field": "as_of_date", "label": "기준일", "required": False,
             "note": "YYYY-MM-DD. 없으면 업로드할 때 지정한 날짜를 씁니다.",
             "aliases": ["기준일", "실사일"]},
            {"field": "memo", "label": "메모", "required": False,
             "aliases": ["메모", "비고"]},
        ],
        "notes": [
            "수량은 목표 잔량입니다. 원장이 없으면 최초 재고(OPENING)로, "
            "있으면 차이만큼 조정(ADJUSTMENT)으로 들어갑니다.",
            "잔량을 덮어쓰지 않습니다. 같은 값이면 아무것도 기록하지 않습니다.",
            "같은 파일을 두 번 올려도 재고가 두 번 늘지 않습니다.",
            "상품 이름으로는 찾지 않습니다. 비슷한 이름을 잘못 맞추면 엉뚱한 "
            "상품의 재고가 바뀌기 때문입니다.",
        ],
    })
