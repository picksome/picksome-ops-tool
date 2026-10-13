"""쿠팡 정산서 파일 적재 (M056 · docs/82).

원본 파일을 fixture 로 쓰지 않는다
------------------------------
실제 WING 정산서에는 구매자명·주문번호·금액이 들어 있다. Git 에 넣으면
되돌릴 수 없다. 대신 **실측한 스키마 그대로** 최소 synthetic 파일을 만든다 —
컬럼 25개와 행 구조(상품 + 기본배송료 + 추가배송료)는 실제 파일에서 확인한
것이고, 값만 지어낸 것이다.

환불은 "지원한다" 고 말하지 않는다
------------------------------
실제 파일 3개(67행)에 환불 표본이 **0건**이었다(환불수량 전부 0, 취소완료일
전부 비어 있음). 아래에 음수·환불수량 fixture 가 있지만 그건 **파서가
무너지지 않는지** 보는 것이지, 쿠팡 환불 정산 구조를 검증한 것이 아니다.
"""

from __future__ import annotations

import uuid
from datetime import date

import pytest
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.exceptions import ValidationError
from app.io.settlement_import import coupang_msf as msf
from app.services import settlement_import as si
from app.services import settlement_matching as sm
from app.tests.fixtures.factories import (
    make_app_user,
    make_integration_account,
    make_organization,
    make_product,
    make_sales_order,
)

# 실측한 헤더 25개. 순서까지 실제 파일과 같다.
HEADERS = [
    "주문번호", "과세유형", "상품 ID", "상품명", "옵션 ID", "옵션명",
    "판매가", "판매수량", "환불수량", "판매액",
    "판매자 할인쿠폰(A+B)", "판매자 할인쿠폰(A.즉시할인)", "판매자 할인쿠폰(B.다운로드)",
    "판매수수료", "서비스이용율(%,VAT별도)", "서비스이용료소급분", "마이샵수수료할인",
    "정산금액", "구매자명", "결제완료일", "배송완료일", "구매확정일",
    "취소완료일", "구매확정(출고)유형", "정산예정일",
]
DUE = "2026-01-22"
PAID = "2025-12-01"


def _row(order_no, option_id, *, gross=0, fee=0, net=0, coupon="",
         qty=0, refund_qty=0, dates=True, buyer="테스트구매자",
         due=DUE, paid=PAID, cancelled=""):
    """실측 컬럼 순서 그대로 한 행."""
    d = paid if dates else ""
    return [
        order_no, "TAX" if str(option_id).isdigit() else "",
        "9999999999" if str(option_id).isdigit() else "",
        "상품명" if str(option_id).isdigit() else "",
        option_id,
        "옵션명" if str(option_id).isdigit() else "",
        str(gross), qty, refund_qty, str(gross),
        coupon, coupon, coupon,
        str(fee), "10.8", "" if coupon == "" else "0", "" if coupon == "" else "0",
        str(net), buyer, d, d, d, cancelled,
        "수동" if str(option_id).isdigit() else "", due if dates else "",
    ]


def _order_rows(order_no, *, gross, fee, net, coupon="0", ship_gross=0,
                ship_fee=0, ship_net=0, due=DUE, paid=PAID):
    """실제 파일의 주문 1건 = 상품 + 기본배송료 + 추가배송료."""
    return [
        _row(order_no, "92195434920", gross=gross, fee=fee, net=net,
             coupon=coupon, qty=1, due=due, paid=paid),
        _row(order_no, "<기본배송료>", gross=ship_gross, fee=ship_fee,
             net=ship_net, dates=ship_gross > 0, due=due, paid=paid),
        _row(order_no, "<추가배송료>", gross=0, fee=0, net=0,
             dates=False, due=due, paid=paid),
    ]


def _write(tmp_path, name, rows, headers=HEADERS):
    wb = Workbook()
    ws = wb.active
    ws.title = "Excel Data (1)"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    p = tmp_path / name
    wb.save(p)
    wb.close()
    return p


@pytest.fixture
def org(db: Session) -> uuid.UUID:
    return make_organization(db)


@pytest.fixture
def actor(db: Session, org: uuid.UUID) -> uuid.UUID:
    return make_app_user(db, org_id=org, role="ADMIN")


@pytest.fixture
def account(db: Session, org: uuid.UUID) -> uuid.UUID:
    """쿠팡 판매계정. 채널을 선언해 둔다 — 서비스가 그걸 검사한다."""
    return make_integration_account(
        db, org_id=org, provider="COUPANG", account_label="쿠팡 테스트 스토어")


def _sale(db, org, order_no, *, price=24000, marketplace="COUPANG"):
    p = make_product(db, org_id=org, cost_amount=1000)
    return make_sales_order(
        db, org_id=org, product_id=p, sale_date=date(2025, 12, 1),
        sale_price=price, marketplace=marketplace, marketplace_order_no=order_no)


def _run(db, org, account, path, mode=si.COMMIT, actor_user_id=None):
    return si.import_file(db, org_id=org, integration_account_id=account,
                          path=path, mode=mode, actor_user_id=actor_user_id)


# ──────────────────────────────────────────────────────────────────────
# 1. 파서 — 실측 구조
# ──────────────────────────────────────────────────────────────────────

def test_row_type_comes_from_option_id(tmp_path):
    """`옵션 ID` 가 행 종류 판별자다. 실측 22개 주문 전부 이 구조였다."""
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    r = msf.parse_workbook(p)
    assert [x.row_type for x in r.rows] == [
        msf.PRODUCT_SALE, msf.SHIPPING_FEE_BASE, msf.SHIPPING_FEE_EXTRA]
    assert r.reconciles()


def test_unknown_marker_is_unknown_not_product(tmp_path):
    """모르는 표식을 상품으로 세지 않는다."""
    p = _write(tmp_path, "a.xlsx", [_row("1001", "<처음보는것>", gross=100, fee=0, net=100)])
    r = msf.parse_workbook(p)
    assert r.rows[0].row_type == msf.UNKNOWN


def test_float_order_no_loses_no_digits(tmp_path):
    """Excel 이 주문번호를 float 로 준다. `str()` 하면 '.0' 이 붙어 영영 안 붙는다."""
    p = _write(tmp_path, "a.xlsx", [_row(4100157762225.0, "92195434920",
                                         gross=1000, fee=0, net=1000, qty=1)])
    r = msf.parse_workbook(p)
    assert r.rows[0].order_no == "4100157762225"


def test_blank_is_null_not_zero(tmp_path):
    """**빈 값은 0 이 아니다.** 배송료 행에는 쿠폰이라는 개념이 없다."""
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    r = msf.parse_workbook(p)
    product, base, extra = r.rows
    assert product.seller_coupon == 0        # 실제로 0원 쿠폰
    assert base.seller_coupon is None        # 해당 없음
    assert extra.seller_coupon is None


def test_zero_amount_row_is_kept(tmp_path):
    """전부 0 인 추가배송료 행도 원본의 사실이다. 버리면 대사가 안 맞는다."""
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    r = msf.parse_workbook(p)
    extra = r.rows[2]
    assert extra.gross_revenue == 0 and extra.net_settlement == 0
    assert r.source_rows == 3 and len(r.rows) == 3


def test_negative_amount_is_preserved(tmp_path):
    """음수를 abs() 하지 않는다. 실제 표본은 없지만 파서가 무너지면 안 된다."""
    p = _write(tmp_path, "a.xlsx", [_row("1001", "92195434920",
                                         gross=-5000, fee=-540, net=-4460, coupon="0", qty=1)])
    r = msf.parse_workbook(p)
    assert r.rows[0].gross_revenue == -5000
    assert r.rows[0].net_settlement == -4460
    assert r.rows[0].is_refund_like


def test_missing_required_header_fails_loud(tmp_path):
    """필수 컬럼이 없으면 반쯤 읽지 않는다."""
    headers = [h for h in HEADERS if h != "정산금액"]
    p = _write(tmp_path, "a.xlsx", [], headers=headers)
    with pytest.raises(msf.UnsupportedSchemaError) as e:
        msf.parse_workbook(p)
    assert "정산금액" in str(e.value)


def test_unknown_optional_header_is_reported_not_fatal(tmp_path):
    """벤더가 컬럼을 늘려도 적재는 돈다. 다만 조용히 넘기지 않는다."""
    headers = HEADERS + ["신규정산항목"]
    rows = [r + ["1234"] for r in _order_rows("1001", gross=100, fee=0, net=100)]
    p = _write(tmp_path, "a.xlsx", rows, headers=headers)
    r = msf.parse_workbook(p)
    assert r.unknown_headers == ["신규정산항목"]
    assert len(r.rows) == 3


def test_duplicate_logical_row_in_one_file_is_rejected(tmp_path):
    """한 파일 안 같은 식별자 두 번. 조용히 합치면 한 행이 사라진다."""
    rows = _order_rows("1001", gross=100, fee=0, net=100)
    p = _write(tmp_path, "a.xlsx", rows + [rows[0]])
    r = msf.parse_workbook(p)
    assert len(r.rows) == 3
    assert [x.reason for x in r.rejected] == ["DUPLICATE_IN_FILE"]
    assert r.reconciles()          # 거부돼도 대사에는 들어간다


def test_unreadable_amount_is_rejected_not_zeroed(tmp_path):
    p = _write(tmp_path, "a.xlsx", [_row("1001", "92195434920", gross=100, fee=0, net=100)])
    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb.active
    ws.cell(row=2, column=HEADERS.index("정산금액") + 1, value="알 수 없음")
    wb.save(p)
    wb.close()
    r = msf.parse_workbook(p)
    assert len(r.rows) == 0
    assert r.rejected[0].reason == "INVALID_VALUE"
    assert r.reconciles()


def test_snapshot_file_shape_is_unsupported(tmp_path):
    """미확정 스냅샷은 정산서가 아니다 — 억지로 읽지 않는다."""
    snap = ["주문번호", "상품번호", "상품명", "구매자명", "판매가", "판매금액",
            "판매배송비", "취소금액", "취소배송비", "주문일", "배송완료일", "배송상태"]
    p = _write(tmp_path, "s.xlsx", [], headers=snap)
    with pytest.raises(msf.UnsupportedSchemaError):
        msf.parse_workbook(p)


def test_idempotency_key_includes_due_date(tmp_path):
    """같은 주문·옵션이 다음 정산 회차에 다시 오면 **다른 행**이다.

    예정일을 빼면 나중 회차의 환불 행이 중복으로 취급돼 조용히 사라진다.
    """
    a = msf.idempotency_key(due_date=date(2026, 1, 22), order_no="1", line_reference="9")
    b = msf.idempotency_key(due_date=date(2026, 2, 25), order_no="1", line_reference="9")
    assert a != b


# ──────────────────────────────────────────────────────────────────────
# 2. 적재 · 매칭
# ──────────────────────────────────────────────────────────────────────

def test_dry_run_writes_nothing(db: Session, org, account, tmp_path):
    _sale(db, org, "1001")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    before = db.execute(text("SELECT count(*) FROM marketplace_settlement_record")).scalar()
    s = _run(db, org, account, p, mode=si.DRY_RUN)
    after = db.execute(text("SELECT count(*) FROM marketplace_settlement_record")).scalar()
    assert s.inserted == 3 and after == before
    assert db.execute(text(
        "SELECT count(*) FROM batch_run WHERE org_id = :o"), {"o": str(org)}).scalar() == 0


def test_single_line_order_matches_all_three_rows(db: Session, org, account, tmp_path):
    """배송료 행도 같은 규칙으로 붙는다 — 주문에 라인이 하나뿐이면 모호하지 않다."""
    order = _sale(db, org, "1001")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    s = _run(db, org, account, p)
    assert (s.matched_exact, s.ambiguous, s.unmatched) == (3, 0, 0)
    rows = db.execute(text("""
        SELECT row_type, matching_status, matched_sales_order_id
          FROM marketplace_settlement_record WHERE org_id = :o ORDER BY row_type
    """), {"o": str(org)}).all()
    assert {r[0] for r in rows} == {
        msf.PRODUCT_SALE, msf.SHIPPING_FEE_BASE, msf.SHIPPING_FEE_EXTRA}
    assert all(r[1] == sm.MATCHED_EXACT and r[2] == order for r in rows)


def test_multi_line_order_is_ambiguous(db: Session, org, account, tmp_path):
    """장부에 라인이 둘이면 **어느 라인의 정산금인지 알 수 없다.**

    정산서에 옵션 ID 가 있어도 우리 쪽에 대응하는 식별자가 없다.
    """
    _sale(db, org, "1001", price=24000)
    _sale(db, org, "1001", price=10000)
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    s = _run(db, org, account, p)
    assert (s.matched_exact, s.ambiguous, s.unmatched) == (0, 3, 0)
    assert db.execute(text("""
        SELECT count(*) FROM marketplace_settlement_record
         WHERE org_id = :o AND matched_sales_order_id IS NOT NULL
    """), {"o": str(org)}).scalar() == 0


def test_unknown_order_is_unmatched(db: Session, org, account, tmp_path):
    p = _write(tmp_path, "a.xlsx", _order_rows("9999", gross=100, fee=0, net=100))
    s = _run(db, org, account, p)
    assert (s.matched_exact, s.ambiguous, s.unmatched) == (0, 0, 3)


def test_reconciliation_holds_with_mixed_outcomes(db: Session, org, account, tmp_path):
    """**한 행도 사라지지 않는다.** 원본 = 적재 + 거부."""
    _sale(db, org, "1001")
    _sale(db, org, "1002")
    _sale(db, org, "1002")                               # multi-line
    rows = (_order_rows("1001", gross=24000, fee=2592, net=21408)
            + _order_rows("1002", gross=10000, fee=1080, net=8920)
            + _order_rows("1003", gross=5000, fee=540, net=4460))
    rows.append(rows[0])                                 # 논리 중복 → 거부
    p = _write(tmp_path, "a.xlsx", rows)
    s = _run(db, org, account, p)
    assert s.source_rows == 10
    assert s.persisted + s.rejected == s.source_rows
    assert s.reconciles()
    assert (s.matched_exact, s.ambiguous, s.unmatched, s.rejected) == (3, 3, 3, 1)


# ──────────────────────────────────────────────────────────────────────
# 3. 멱등성
# ──────────────────────────────────────────────────────────────────────

def test_same_file_twice_inserts_nothing_new(db: Session, org, account, tmp_path):
    _sale(db, org, "1001")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    first = _run(db, org, account, p)
    second = _run(db, org, account, p)
    assert first.inserted == 3
    assert second.inserted == 0 and second.unchanged == 3
    assert db.execute(text(
        "SELECT count(*) FROM marketplace_settlement_record WHERE org_id = :o"),
        {"o": str(org)}).scalar() == 3


def test_same_logical_row_in_a_different_file_is_not_duplicated(
        db: Session, org, account, tmp_path):
    """**파일 해시로 멱등성을 판단하지 않는다.**

    같은 기간을 다시 내려받으면 파일이 달라도 논리적으로 같은 행이다.
    """
    _sale(db, org, "1001")
    rows = _order_rows("1001", gross=24000, fee=2592, net=21408)
    a = _write(tmp_path, "a.xlsx", rows)
    b = _write(tmp_path, "b.xlsx", list(reversed(rows)))   # 순서만 다르다
    assert msf.file_sha256(a) != msf.file_sha256(b)
    _run(db, org, account, a)
    second = _run(db, org, account, b)
    assert second.inserted == 0 and second.unchanged == 3


def test_changed_amount_is_surfaced_not_overwritten(db: Session, org, account, tmp_path):
    """재발행으로 금액이 달라졌으면 **덮어쓰지 않고 보여 준다.**

    그건 새 사실이고, 정정은 settlement_adjustment 의 일이다.
    """
    _sale(db, org, "1001")
    a = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    _run(db, org, account, a)
    b = _write(tmp_path, "b.xlsx", _order_rows("1001", gross=24000, fee=3000, net=21000))
    s = _run(db, org, account, b)
    assert s.inserted == 0
    assert s.changed_not_applied == 1        # 상품 행만 금액이 바뀌었다
    assert s.unchanged == 2
    kept = db.execute(text("""
        SELECT marketplace_fee FROM marketplace_settlement_record
         WHERE org_id = :o AND row_type = 'PRODUCT_SALE'
    """), {"o": str(org)}).scalar()
    assert kept == 2592                      # 원본이 그대로 남는다


def test_different_settlement_cycles_coexist(db: Session, org, account, tmp_path):
    """같은 주문·옵션이 다음 회차에 다시 와도 **다른 행**이다."""
    _sale(db, org, "1001")
    a = _write(tmp_path, "a.xlsx",
               _order_rows("1001", gross=24000, fee=2592, net=21408, due="2026-01-22"))
    b = _write(tmp_path, "b.xlsx",
               _order_rows("1001", gross=-24000, fee=-2592, net=-21408, due="2026-02-25"))
    _run(db, org, account, a)
    s = _run(db, org, account, b)
    assert s.inserted == 3
    assert db.execute(text(
        "SELECT count(*) FROM marketplace_settlement_record WHERE org_id = :o"),
        {"o": str(org)}).scalar() == 6


# ──────────────────────────────────────────────────────────────────────
# 4. 저장한 것 / 저장하지 않은 것
# ──────────────────────────────────────────────────────────────────────

def test_buyer_name_is_never_persisted(db: Session, org, account, tmp_path):
    """개인정보는 담을 그릇 자체가 없다."""
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=100, fee=0, net=100))
    r = msf.parse_workbook(p)
    assert not any("구매자" in f for f in r.rows[0].__slots__)
    _sale(db, org, "1001")
    _run(db, org, account, p)
    # 어느 텍스트 컬럼에도 구매자명이 없다.
    hit = db.execute(text("""
        SELECT count(*) FROM marketplace_settlement_record
         WHERE org_id = :o
           AND (COALESCE(marketplace_order_no,'') LIKE '%테스트구매자%'
             OR COALESCE(line_reference,'') LIKE '%테스트구매자%'
             OR COALESCE(match_reason,'') LIKE '%테스트구매자%')
    """), {"o": str(org)}).scalar()
    assert hit == 0


def test_expected_date_is_not_settled_at(db: Session, org, account, tmp_path):
    """정산예정일을 `settled_at` 에 넣지 않는다 — 입금 증거가 아니다."""
    _sale(db, org, "1001")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    _run(db, org, account, p)
    row = db.execute(text("""
        SELECT settled_at, expected_settlement_date, settlement_period_start,
               settlement_period_end
          FROM marketplace_settlement_record
         WHERE org_id = :o AND row_type = 'PRODUCT_SALE'
    """), {"o": str(org)}).first()
    assert row[0] is None                       # settled_at
    assert row[1] == date(2026, 1, 22)          # expected
    assert row[2] == date(2025, 12, 1)          # 기간은 결제완료일 기준
    assert row[3] == date(2025, 12, 1)


def test_provenance_without_raw_dump(db: Session, org, account, tmp_path):
    """원본 행을 통째로 저장하지 않는다. 해시·지문만 남는다."""
    _sale(db, org, "1001")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    _run(db, org, account, p)
    row = db.execute(text("""
        SELECT import_file_hash, source_row_fingerprint, import_idempotency_key,
               raw_json, import_batch_id, line_reference
          FROM marketplace_settlement_record
         WHERE org_id = :o AND row_type = 'PRODUCT_SALE'
    """), {"o": str(org)}).mappings().first()
    assert len(row["import_file_hash"]) == 64
    assert len(row["source_row_fingerprint"]) == 64
    assert row["import_idempotency_key"].startswith(msf.SOURCE_DOCUMENT)
    assert row["raw_json"] is None              # **통째 저장 안 함**
    assert row["import_batch_id"] is not None
    assert row["line_reference"] == "92195434920"


def test_batch_run_is_org_scoped(db: Session, org, account, tmp_path):
    _sale(db, org, "1001")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=100, fee=0, net=100))
    s = _run(db, org, account, p)
    got = db.execute(text("SELECT org_id, status, affected_rows FROM batch_run WHERE id = :i"),
                     {"i": str(s.batch_run_id)}).first()
    assert got[0] == org and got[1] == "SUCCESS" and got[2] == 3


# ──────────────────────────────────────────────────────────────────────
# 5. 격리 · 계정
# ──────────────────────────────────────────────────────────────────────

def test_other_org_order_is_never_matched(db: Session, org, account, tmp_path):
    """남의 조직 주문에 붙지 않는다."""
    other = make_organization(db)
    _sale(db, other, "1001")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=100, fee=0, net=100))
    s = _run(db, org, account, p)
    assert s.unmatched == 3 and s.matched_exact == 0


def test_account_from_other_org_is_rejected(db: Session, org, tmp_path):
    other = make_organization(db)
    foreign = make_integration_account(
        db, org_id=other, provider="COUPANG", account_label="남의 스토어")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=100, fee=0, net=100))
    from app.core.exceptions import NotFoundError
    with pytest.raises(NotFoundError):
        si.import_file(db, org_id=org, integration_account_id=foreign,
                       path=p, mode=si.DRY_RUN)


def test_non_coupang_account_is_rejected(db: Session, org, tmp_path):
    """쿠팡 정산서를 네이버 계정에 넣으면 스토어별 손익이 통째로 틀어진다."""
    aid = make_integration_account(
        db, org_id=org, provider="SMARTSTORE", account_label="네이버 스토어")
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=100, fee=0, net=100))
    with pytest.raises(ValidationError):
        si.import_file(db, org_id=org, integration_account_id=aid,
                       path=p, mode=si.DRY_RUN)


def test_import_does_not_touch_sales_or_cash(db: Session, org, account, tmp_path):
    """정산서를 읽었다고 판매·원가·현금·마감을 건드리지 않는다."""
    order = _sale(db, org, "1001")
    before = db.execute(text("""
        SELECT sale_price, purchase_cost, updated_at FROM sales_order WHERE id = :i
    """), {"i": str(order)}).first()
    cash_before = db.execute(text("SELECT count(*) FROM manual_cash_event")).scalar()
    p = _write(tmp_path, "a.xlsx", _order_rows("1001", gross=24000, fee=2592, net=21408))
    _run(db, org, account, p)
    after = db.execute(text("""
        SELECT sale_price, purchase_cost, updated_at FROM sales_order WHERE id = :i
    """), {"i": str(order)}).first()
    assert before == after
    assert db.execute(text("SELECT count(*) FROM manual_cash_event")).scalar() == cash_before
    assert db.execute(text("SELECT count(*) FROM settlement_adjustment")).scalar() == 0


def test_batch_lookup_does_not_do_n_plus_one(db: Session, org, account, tmp_path):
    """주문 조회는 행 수와 무관하게 **한 번**이다."""
    rows = []
    for i in range(30):
        _sale(db, org, f"200{i:03d}")
        rows += _order_rows(f"200{i:03d}", gross=1000, fee=100, net=900)
    p = _write(tmp_path, "big.xlsx", rows)

    seen: list[str] = []
    from sqlalchemy import event
    engine = db.get_bind()

    def before(conn, cur, statement, params, ctx, many):
        if "FROM sales_order so" in statement:
            seen.append(statement)

    event.listen(engine, "before_cursor_execute", before)
    try:
        s = _run(db, org, account, p, mode=si.DRY_RUN)
    finally:
        event.remove(engine, "before_cursor_execute", before)
    assert s.source_rows == 90
    assert len(seen) == 1, f"주문 조회가 {len(seen)}번 돌았다 (N+1)"


# ──────────────────────────────────────────────────────────────────────
# 6. 운영 화면이 셋을 뭉치지 않는다
# ──────────────────────────────────────────────────────────────────────

def test_exception_center_splits_ambiguous_from_unmatched(
        db: Session, org, account, tmp_path):
    """예외 센터가 **다른 항목**으로 낸다 — 할 일이 다르기 때문이다."""
    from app.services import exception_center

    _sale(db, org, "1001")
    _sale(db, org, "1001")                       # multi-line → AMBIGUOUS
    rows = (_order_rows("1001", gross=24000, fee=2592, net=21408)
            + _order_rows("2002", gross=5000, fee=540, net=4460))  # → UNMATCHED
    _run(db, org, account, _write(tmp_path, "a.xlsx", rows))

    kinds = {
        i["kind"]: i["title"]
        for i in exception_center.build(db, org_id=org)["items"]
        if i["category"] == "SETTLEMENT"
    }
    assert "AMBIGUOUS_SETTLEMENT" in kinds
    assert "UNMATCHED_SETTLEMENT" in kinds
    assert kinds["AMBIGUOUS_SETTLEMENT"] != kinds["UNMATCHED_SETTLEMENT"]


def test_excluded_rows_drop_out_of_exception_center(
        db: Session, org, account, tmp_path):
    """제외한 행을 계속 세면 숫자가 영원히 줄지 않는다."""
    from app.services import exception_center

    rows = _order_rows("9999", gross=5000, fee=540, net=4460)
    _run(db, org, account, _write(tmp_path, "a.xlsx", rows))

    def unmatched_titles():
        return [i["title"] for i in exception_center.build(db, org_id=org)["items"]
                if i["kind"] == "UNMATCHED_SETTLEMENT"]

    assert unmatched_titles(), "미매칭이 안 잡히면 이 테스트가 무의미하다"

    ids = db.execute(text("""
        SELECT id FROM marketplace_settlement_record
         WHERE org_id = :o AND matching_status = 'UNMATCHED'
    """), {"o": str(org)}).scalars().all()
    for rid in ids:
        sm.mark_excluded(db, org_id=org, record_id=rid,
                         reason="주문과 무관한 정산행", actor_user_id=None)

    assert unmatched_titles() == []
